"""Fill slide 2 (The Problem) for Udyog-drishti AI of the Seva Innovation Challenge 2026 template.

Run:  python build_slide2.py   (needs python-pptx)
Reads template_original.pptx and writes Udyog-drishti_AI_Seva_2026.pptx next to it.
"""

import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
SRC, OUT = HERE / "template_original.pptx", HERE / "Udyog-drishti_AI_Seva_2026.pptx"

INK = RGBColor(0x1F, 0x29, 0x37)
NAVY = RGBColor(0x1B, 0x2A, 0x5C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def fill(shape, paras, color=INK, size=10, bullets=False, space_after=3):
    """paras: list of paragraphs; each is a str or a list of (text, bold) runs."""
    tf = shape.text_frame
    first = tf.paragraphs[0]
    proto_ppr = copy.deepcopy(first._p.pPr) if first._p.pPr is not None else None
    proto_rpr = first.runs[0]._r.rPr if first.runs else None
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, para in enumerate(paras):
        p = first if i == 0 else tf.add_paragraph()
        if i and proto_ppr is not None:
            p._p.insert(0, copy.deepcopy(proto_ppr))
        runs = [(para, False)] if isinstance(para, str) else para
        for text, bold in runs:
            r = p.add_run()
            if proto_rpr is not None:
                r._r.insert(0, copy.deepcopy(proto_rpr))
            r.text = text
            f = r.font
            f.italic = False
            f.bold = bold
            f.size = Pt(size)
            f.color.rgb = color
        p.space_after = Pt(space_after)
        if bullets:
            ppr = p._p.get_or_add_pPr()
            ppr.set("marL", str(Inches(0.14)))
            ppr.set("indent", str(-Inches(0.14)))
            for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
                for el in ppr.findall(qn(tag)):
                    ppr.remove(el)
            bu = etree.SubElement(ppr, qn("a:buChar"))
            bu.set("char", "•")


prs = Presentation(SRC)
s = prs.slides[1]
by_name = {sh.name: sh for sh in s.shapes}

fill(by_name["Text 7"], [[
    ("Big industries are checked mainly by stack monitors ", False), ("they run themselves", True),
    (" and by occasional inspections. Since July 2025 most coal plants no longer need SO₂", False),
    (" scrubbers, and no one independently tracks what each plant and industrial cluster releases, month by month.", False)]])

fill(by_name["Text 10"], [[
    ("People living near Maharashtra's coal plants and ", False), ("critically polluted clusters", True),
    (": Chandrapur, Koradi–Khaparkheda (Nagpur), Tarapur, Taloja, Dombivli. Also MPCB, CPCB and NGT committees, "
     "who need independent evidence to act.", False)]])

fill(by_name["Text 13"], [
    [("Koradi, Khaparkheda, Tiroda, Mauda", True), (": 9,360 MW exempted from SO₂ scrubbers (Jul 2025)", False)],
    [("Chandrapur and Koradi", True), (": among India's top SO₂ hotspots in NASA data", False)],
    [("Nagpur over the PM10 limit on ", False), ("46 of 92", True), (" spring days in 2026", False)],
], size=9.5, bullets=True)

fill(by_name["Text 16"], [
    [("Self-reported monitors", True), (" (OCEMS): run by the plant, with gaps and poor data quality", False)],
    [("Inspections are occasional", True), ("; no picture between visits", False)],
    [("No independent check", True), (" that scrubbers or effluent plants actually cut pollution", False)],
], size=9.5, bullets=True)

fill(by_name["Text 18"], [[("78%", True)]], color=RGBColor(0xF2, 0x8C, 0x28), size=28)
fill(by_name["Text 19"], ["of India's coal power units no longer need SO₂ scrubbers (July 2025)"], color=WHITE, size=10.5)
fill(by_name["Text 21"], ["Sulphur and ash keep reaching lungs, farms and rivers around Nagpur and Chandrapur, "
                          "and the only proof of compliance is the polluter's own data."],
     color=RGBColor(0xC9, 0xD2, 0xEC), size=10)

src = s.shapes.add_textbox(Inches(0.5), Inches(5.1), Inches(5.4), Inches(0.38))
src.name = "Sources"
tf = src.text_frame
tf.word_wrap = True
for side in ("left", "right", "top", "bottom"):
    setattr(tf, f"margin_{side}", 0)
r = tf.paragraphs[0].add_run()
r.text = ("Sources: MoEFCC notification, 11 Jul 2025; Greenpeace analysis of NASA OMI data; TechTimes analysis of "
          "CPCB data, Aug 2026; CEEW 2024 and CSE on OCEMS; CREA, Jul 2026.")
r.font.size = Pt(7)
r.font.name = "Nirmala UI"
r.font.color.rgb = MUTED

prs.save(OUT)
print("saved", OUT)
