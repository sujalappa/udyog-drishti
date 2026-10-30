"""Fill slide 1 (title) of the Seva Innovation Challenge 2026 deck.

Run after build_slide2.py and build_slides_3_6.py; edits the deck in place.
"""

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
DEMO_LINK = "udyog-drishti.onrender.com"  # live demo on Render

CONFIG = {
    "deck": "Udyog-drishti_AI_Seva_2026.pptx",
    "sector": "SUSTAINABILITY - ENERGY, CLIMATE AND OTHERS",
    "title": "Udyog-drishti AI",
    "line": "A monthly satellite check of SO₂ and NO₂ at Maharashtra's large coal power plants and industrial areas: "
            "where to inspect first, and whether a scrubber worked. Independent of the plants' own monitors.",
    "subtheme": "Sub-theme: Industrial decarbonisation, pollution control and wastewater reuse",
    "image": "img/app_state.png",
    "facts": ["17 plants and industrial areas, monthly since 2019, in ~1 min",
              "Sentinel-5P SO₂ and NO₂; agrees with plants' official filings"],
}

ORANGE = RGBColor(0xF2, 0x8C, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xC9, 0xD2, 0xEC)
FONT = "Nirmala UI"


def set_text(shape, runs, size, color, bold=False, italic=False):
    """Replace a text box's text with runs [(text, bold, color|None)], keeping the first run's formatting."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    proto = copy.deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    for text, b, col in runs:
        r = p.add_run()
        if proto is not None:
            r._r.insert(0, copy.deepcopy(proto))
        r.text = text
        r.font.size, r.font.bold, r.font.italic = Pt(size), b if b is not None else bold, italic
        r.font.color.rgb = col or color


def textbox(slide, x, y, w, h, paras, size, color, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", 0)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        for text, b, col in runs:
            r = p.add_run()
            r.text = text
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), b, col or color
    return tb


def link_demo(prs):
    """Make every occurrence of the demo link clickable (slides 1, 4 and 6)."""
    if DEMO_LINK.startswith("["):
        return
    # PowerPoint draws links in the theme's hyperlink colour (blue); use the deck's orange instead.
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(theme.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for tag in ("hlink", "folHlink"):
        for el in root.findall(f".//a:clrScheme/a:{tag}", ns):
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, f"{{{ns['a']}}}srgbClr", val="F28C28")
    theme._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        if r.text.strip() == DEMO_LINK:
                            r.hyperlink.address = "https://" + DEMO_LINK


def main(cfg=CONFIG):
    deck = HERE / cfg["deck"]
    prs = Presentation(deck)
    s = prs.slides[0]
    sh = {x.name: x for x in s.shapes}
    set_text(sh["Text 0"], [(f"SEVA INNOVATION CHALLENGE 2026  •  {cfg['sector']}", None, None)], 10, ORANGE)
    set_text(sh["Text 1"], [(cfg["title"], True, None)], 36, WHITE)
    set_text(sh["Text 2"], [(cfg["line"], False, None)], 13, SOFT, italic=False)
    sh["Text 2"].height = Inches(0.9)
    textbox(s, 0.5, 3.0, 5.4, 0.3, [[(cfg["subtheme"], True, ORANGE)]], 10.5, ORANGE)
    textbox(s, 0.5, 4.05, 5.6, 0.75, [
        [("Sujal Chondekar  ·  Kshitij Mehrotra", True, WHITE)],
        [("Mentor: Dr. Swati Hira, Assistant Professor, CSE", False, SOFT)],
        [("Indian Institute of Information Technology, Nagpur", False, SOFT)],
    ], 10.5, WHITE)

    # Right card: replace the "How to use this template" guidance with a prototype card.
    card = sh["Shape 15"]
    set_text(sh["Text 16"], [("Working prototype", True, None)], 13, ORANGE)
    sh["Text 17"]._element.getparent().remove(sh["Text 17"]._element)
    x, y, w = card.left / 914400 + 0.2, 1.4, card.width / 914400 - 0.4
    pic = s.shapes.add_picture(str(HERE / cfg["image"]), Inches(x), Inches(y), width=Inches(w))
    pic.line.color.rgb = RGBColor(0x3A, 0x4A, 0x7A)
    top = y + pic.height / 914400 + 0.15
    textbox(s, x, top, w, 1.2, [[("• " + f, False, WHITE)] for f in cfg["facts"]] +
            [[("Live demo: ", True, WHITE), (DEMO_LINK, True, ORANGE)]], 10, WHITE)
    link_demo(prs)
    prs.save(deck)
    print("saved", deck)


if __name__ == "__main__":
    main()
